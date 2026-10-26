"""Builds sotif_validation_targets.xlsx — every Lesson 1 calculation as live Excel formulas."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment

OUT = "/home/claude/sotif-risk/tools/sotif_validation_targets.xlsx"
F = "Arial"
BLUE = Font(name=F, color="0000FF")
BLACK = Font(name=F)
BOLD = Font(name=F, bold=True)
HEAD = Font(name=F, bold=True, color="FFFFFF")
TITLE = Font(name=F, bold=True, size=14)
GREEN = Font(name=F, color="008000")
YELLOW = PatternFill("solid", fgColor="FFFF00")
HEADFILL = PatternFill("solid", fgColor="0E6B63")
thin = Side(style="thin", color="BFBFBF")
BOX = Border(bottom=thin)
SCI = "0.00E+00"
INT = "#,##0"
DEC3 = "0.000"

wb = Workbook()


def title(ws, text, sub):
    ws["A1"] = text; ws["A1"].font = TITLE
    ws["A2"] = sub; ws["A2"].font = Font(name=F, italic=True, color="595959")


def header(ws, row, labels):
    for i, l in enumerate(labels):
        c = ws.cell(row=row, column=1 + i, value=l)
        c.font = HEAD; c.fill = HEADFILL; c.alignment = Alignment(wrap_text=True, vertical="center")


def put(ws, ref, val, font=BLACK, fmt=None, fill=None, note=None):
    c = ws[ref]; c.value = val; c.font = font
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    if note: c.comment = Comment(note, "Course")
    c.border = BOX
    return c


# ---------------- Inputs ----------------
ws = wb.active; ws.title = "Inputs"
title(ws, "Inputs", "Blue text on yellow = edit these. Everything else in the workbook recalculates from them.")
header(ws, 4, ["Parameter", "Symbol", "Value", "Unit", "Source / note"])
rows = [
    ("Acceptance criterion for harm H", "A_H", 1e-7, "1/h", "Example value from Lesson 1, Exercise 1 (not a MOTOR AI figure)."),
    ("Exposure: P(hazardous scenario | hazardous behaviour)", "P(E|HB)", 0.1, "-", "Example. Must be justified for your ODD and routing (Putze et al. 2023)."),
    ("Non-controllability: P(not controlled | E)", "P(C̄|E)", 0.5, "-", "Example. For L4, depends on fallback / minimal-risk manoeuvre."),
    ("Severity: P(severity S reached | not controlled)", "P(S|C̄)", 0.2, "-", "Example. Correct conditioning is P(S|C̄,E,HB) (Putze et al. 2023)."),
    ("Confidence level", "C", 0.95, "-", "Frequentist confidence. In Betschinske et al. 2025, alpha = 1 - C."),
    ("Hazardous behaviours allowed in the test plan", "k", 0, "events", "Must be fixed BEFORE testing; changing it after seeing data invalidates the test."),
    ("Test fleet size", "N_fleet", 50, "vehicles", "Example."),
    ("Operating hours per vehicle per year", "h_yr", 3000, "h/yr", "Example. 24/7 operation = 8766 h/yr."),
]
for i, (p, s, v, u, n) in enumerate(rows):
    r = 5 + i
    put(ws, f"A{r}", p); put(ws, f"B{r}", s)
    put(ws, f"C{r}", v, BLUE, SCI if v < 1e-3 else None, YELLOW)
    put(ws, f"D{r}", u); put(ws, f"E{r}", n)
ws["A14"] = "Legend: blue = input, black = formula, green = link to another sheet."; ws["A14"].font = Font(name=F, italic=True)
for col, w in zip("ABCDE", (50, 12, 14, 10, 80)):
    ws.column_dimensions[col].width = w

# ---------------- Validation Target ----------------
vt = wb.create_sheet("Validation Target")
title(vt, "Acceptance criterion → validation target", "ISO 21448 Annex C.2 decomposition + Poisson test (Lesson 1, Eq. 1–3)")
header(vt, 4, ["Quantity", "Symbol", "Value", "Unit", "Formula / equation"])
calc = [
    ("Acceptance criterion", "A_H", "=Inputs!C5", "1/h", "Input", GREEN, SCI),
    ("Product of conditional factors", "P(E|HB)·P(C̄|E)·P(S|C̄)", "=Inputs!C6*Inputs!C7*Inputs!C8", "-", "Eq. 1 denominator", BLACK, DEC3),
    ("Tolerable hazardous-behaviour rate", "R_HB", "=C5/C6", "1/h", "Eq. 1: R_HB = A_H / [P(E|HB)·P(C̄|E)·P(S|C̄)]", BLACK, SCI),
    ("Confidence", "C", "=Inputs!C9", "-", "Input", GREEN, "0.00"),
    ("Events allowed", "k", "=Inputs!C10", "events", "Input", GREEN, "0"),
    ("Poisson factor", "λ_k(C)", "=CHIINV(1-C8,2*(C9+1))/2", "-", "Eq. 3: λ_k(C) = ½·χ²_{2(k+1)}(C)  (CHIINV takes the right-tail probability 1−C)", BLACK, "0.0000"),
    ("Zero-failure shortcut (k=0)", "−ln(1−C)", "=-LN(1-C8)", "-", "Eq. 2: equals λ_0(C); ≈ 3.00 at C=0.95 (rule of three)", BLACK, "0.0000"),
    ("Required test exposure", "τ_k", "=C10/C7", "h", "Eq. 3: τ_k = λ_k(C) / R_HB", BLACK, INT),
    ("Fleet hours per year", "N_fleet·h_yr", "=Inputs!C11*Inputs!C12", "h/yr", "Inputs", BLACK, INT),
    ("Calendar time with this fleet", "τ_k / (N·h_yr)", "=C12/C13", "years", "", BLACK, "0.00"),
    ("Check: P(X ≤ k | λ_k)", "should equal 1−C", "=POISSON(C9,C10,TRUE)", "-", "Poisson CDF at the plan limit; must equal 1 − C", BLACK, "0.0000"),
    ("Check: 1 − C", "", "=1-C8", "-", "", BLACK, "0.0000"),
    ("Cost of allowing k events vs k=0", "λ_k / λ_0", "=C10/C11", "×", "", BLACK, "0.00"),
]
for i, (q, s, f, u, n, font, fmt) in enumerate(calc):
    r = 5 + i
    put(vt, f"A{r}", q); put(vt, f"B{r}", s); put(vt, f"C{r}", f, font, fmt); put(vt, f"D{r}", u); put(vt, f"E{r}", n)
for col, w in zip("ABCDE", (40, 24, 16, 8, 80)):
    vt.column_dimensions[col].width = w

# ---------------- Chi2 table ----------------
ct = wb.create_sheet("Factor Table")
title(ct, "λ_k(C) = ½·χ²_{2(k+1)}(C)", "Exposure multiplier for plans allowing k events. τ_k = λ_k(C) / R_HB.")
ct["A4"] = "C \\ k"; ct["A4"].font = HEAD; ct["A4"].fill = HEADFILL
for j in range(6):
    c = ct.cell(row=4, column=2 + j, value=j); c.font = HEAD; c.fill = HEADFILL
for i, C in enumerate((0.90, 0.95, 0.99)):
    r = 5 + i
    put(ct, f"A{r}", C, BLUE, "0.00", YELLOW)
    for j in range(6):
        col = chr(ord("B") + j)
        put(ct, f"{col}{r}", f"=CHIINV(1-$A{r},2*({col}$4+1))/2", BLACK, DEC3)
ct["A9"] = "Ratio vs k = 0"; ct["A9"].font = BOLD
for i in range(3):
    r = 10 + i
    put(ct, f"A{r}", f"=A{5+i}", BLACK, "0.00")
    for j in range(6):
        col = chr(ord("B") + j)
        put(ct, f"{col}{r}", f"={col}{5+i}/$B{5+i}", BLACK, "0.00")
ct["A14"] = "Price of confidence (k = 0)"; ct["A14"].font = BOLD
put(ct, "A15", "τ(0.99) / τ(0.95)"); put(ct, "B15", "=LN(1-0.99)/LN(1-0.95)", BLACK, "0.000")
put(ct, "A16", "τ(0.95) / τ(0.90)"); put(ct, "B16", "=LN(1-0.95)/LN(1-0.90)", BLACK, "0.000")
ct.column_dimensions["A"].width = 28
for col in "BCDEFG":
    ct.column_dimensions[col].width = 11

# ---------------- Benchmarks ----------------
bm = wb.create_sheet("Benchmarks")
title(bm, "Benchmarks", "Kalra & Paddock (RAND RR-1478, 2016) and Schaller et al. (Discover AI 6:1172, 2026)")
header(bm, 4, ["Quantity", "Value", "Unit", "Formula / source"])
b = [
    ("RAND: US fatality rate (2013)", 1.09e-8, "1/mile", "Kalra & Paddock 2016: 1.09 fatalities per 100 million miles (32,719 fatalities)", BLUE, SCI),
    ("Confidence", 0.95, "-", "Kalra & Paddock 2016", BLUE, "0.00"),
    ("Failure-free miles, exact", "=LN(1-B6)/LN(1-B5)", "miles", "n = ln(1−C)/ln(1−F)  (RAND Eq. 2, R = 1−F)", BLACK, INT),
    ("Failure-free miles, Poisson approx.", "=-LN(1-B6)/B5", "miles", "n ≈ −ln(1−C)/F   (Lesson 1, Eq. 2)", BLACK, INT),
    ("Fleet: vehicles", 100, "vehicles", "RAND scenario", BLUE, "0"),
    ("Fleet: average speed", 25, "mph", "RAND scenario, 24/7 operation", BLUE, "0"),
    ("Years to drive the exact miles", "=B7/(B9*B10*24*365)", "years", "RAND: ≈ 12.5 years", BLACK, "0.0"),
    ("Schaller: example criterion", 1e-9, "1/h", "Schaller et al. 2026 §5.1: overlooking a car on the other lane ≤ 1e-9 per hour", BLUE, SCI),
    ("Failure-free hours at C (B6)", "=-LN(1-B6)/B12", "h", "τ = −ln(1−C)/R", BLACK, SCI),
    ("Hours per year, 24/7", 8766, "h/yr", "365.25 × 24", BLUE, INT),
    ("Single-vehicle years", "=B13/B14", "years", "", BLACK, INT),
    ("Fleet size for comparison", 10000, "vehicles", "Illustrative", BLUE, INT),
    ("Fleet years", "=B15/B16", "years", "", BLACK, "0.0"),
]
for i, (q, v, u, n, font, fmt) in enumerate(b):
    r = 5 + i
    put(bm, f"A{r}", q); put(bm, f"B{r}", v, font, fmt, YELLOW if font is BLUE else None); put(bm, f"C{r}", u); put(bm, f"D{r}", n)
for col, w in zip("ABCD", (40, 18, 10, 85)):
    bm.column_dimensions[col].width = w

# ---------------- Exercises ----------------
ex = wb.create_sheet("Exercises")
title(ex, "Lesson 1 exercises — worked", "Independent of the Inputs sheet so the answers stay fixed.")
header(ex, 4, ["Item", "Value", "Unit", "Formula"])
e = [
    ("E1/E2: A_H", 1e-7, "1/h", "Given", BLUE, SCI),
    ("P(E|HB)", 0.1, "-", "Given", BLUE, "0.00"),
    ("P(C̄|E)", 0.5, "-", "Given", BLUE, "0.00"),
    ("P(S|C̄)", 0.2, "-", "Given", BLUE, "0.00"),
    ("C", 0.95, "-", "Given", BLUE, "0.00"),
    ("R_HB", "=B5/(B6*B7*B8)", "1/h", "Eq. 1", BLACK, SCI),
    ("E1 answer: τ (k = 0)", "=-LN(1-B9)/B10", "h", "Eq. 2 → 299,573 h", BLACK, INT),
    ("E2 answer: τ₁ (k = 1)", "=CHIINV(1-B9,4)/2/B10", "h", "Eq. 3 → ≈ 474,386 h", BLACK, INT),
    ("E2: cost vs k = 0", "=B12/B11", "×", "≈ 1.58", BLACK, "0.00"),
    ("E3 answer: τ(0.99)/τ(0.95)", "=LN(1-0.99)/LN(1-0.95)", "×", "≈ 1.54", BLACK, "0.000"),
]
for i, (q, v, u, n, font, fmt) in enumerate(e):
    r = 5 + i
    put(ex, f"A{r}", q); put(ex, f"B{r}", v, font, fmt, YELLOW if font is BLUE else None); put(ex, f"C{r}", u); put(ex, f"D{r}", n)
for col, w in zip("ABCD", (32, 16, 8, 40)):
    ex.column_dimensions[col].width = w

# ---------------- Sensitivity ----------------
se = wb.create_sheet("Sensitivity")
title(se, "Required failure-free hours τ = −ln(1−C) / R_HB", "Rows: tolerable rate R_HB. Columns: confidence C. Shows τ ∝ 1/R but only ∝ −ln(1−C).")
se["A4"] = "R_HB (1/h) \\ C"; se["A4"].font = HEAD; se["A4"].fill = HEADFILL
confs = (0.5, 0.9, 0.95, 0.99, 0.999)
for j, C in enumerate(confs):
    c = put(se, f"{chr(66+j)}4", C, BLUE, "0.000", YELLOW)
for i, ex_ in enumerate(range(-4, -10, -1)):
    r = 5 + i
    put(se, f"A{r}", 10.0 ** ex_, BLUE, SCI, YELLOW)
    for j in range(len(confs)):
        col = chr(66 + j)
        put(se, f"{col}{r}", f"=-LN(1-{col}$4)/$A{r}", BLACK, SCI)
se.column_dimensions["A"].width = 18
for col in "BCDEF":
    se.column_dimensions[col].width = 13

# ---------------- Assumptions & Sources ----------------
src = wb.create_sheet("Assumptions & Sources")
title(src, "What a passed test silently assumes", "Ask these of any validation-target result.")
items = [
    "1. Homogeneous Poisson process: constant rate, independent events.",
    "2. Representative exposure: test mix of conditions = deployed mix (per ODD stratum).",
    "3. Frozen software: all hours on the release candidate, or an argued carry-over.",
    "4. One behaviour per harm: if several HBs lead to H, rates add — allocate the budget (Bonferroni, Putze et al. Eq. 2).",
    "5. Frequentist meaning: a system exactly at R passes with probability ≤ 1−C; not P(rate < R) = C without a stated prior.",
    "6. Factors in Eq. 1 are system-dependent (route, ODD, fallback) and need their own justification.",
    "7. k and τ are fixed before testing; extending after an observed event inflates the type-I error.",
    "",
    "Sources",
    "ISO 21448:2022, Clauses 6 and 9, Annex C.2–C.3 (licensed standard).",
    "Putze, Westhofen, Koopmann, Böde, Neurohr (2023). On Quantification for SOTIF Validation of Automated Driving Systems. arXiv:2304.10170",
    "Betschinske et al. (2025). Towards more efficient quantitative safety validation of residual risk for assisted and automated driving. arXiv:2506.10363",
    "Kalra, N. & Paddock, S. M. (2016). Driving to Safety. RAND RR-1478. https://www.rand.org/pubs/research_reports/RR1478.html",
    "Schaller et al. (2026). Addressing SOTIF with a runtime monitor ensemble for AI-based perception systems. Discover AI 6:1172. doi:10.1007/s44163-026-02186-z",
]
for i, t in enumerate(items):
    c = src.cell(row=4 + i, column=1, value=t); c.font = BOLD if t == "Sources" else BLACK
src.column_dimensions["A"].width = 140

for s in wb.worksheets:
    s.freeze_panes = "A5"
wb.save(OUT)
print("saved", OUT)
